"""
excel_export.py — Формирование Excel-отчёта по результатам парсинга реестра договоров.

Структура:
  - Лист «Договоры» — все договоры по всем БИН, каждая строка = один договор
  - Лист «Сводка» — агрегированные итоги по каждому БИН + индикатор загрузки
"""

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
    Border,
    Side,
)
from openpyxl.utils import get_column_letter

from scraper import ScrapeResult, ContractRecord

# ── Стили ──────────────────────────────────────────────────────────────────

HEADER_FONT   = Font(name="Arial", bold=True, size=10, color="000000")
CELL_FONT     = Font(name="Arial", size=10)
LINK_FONT     = Font(name="Arial", size=10, color="0563C1", underline="single")
SUMMARY_FONT  = Font(name="Arial", bold=True, size=10)
HEADER_FILL   = PatternFill(fill_type="solid", fgColor="D9D9D9")
SUMMARY_FILL  = PatternFill(fill_type="solid", fgColor="F2F2F2")
TOTAL_FILL    = PatternFill(fill_type="solid", fgColor="BDD7EE")
RED_FONT      = Font(name="Arial", size=10, color="C00000")
RED_BOLD_FONT = Font(name="Arial", bold=True, size=10, color="C00000")
GREEN_FILL    = PatternFill(fill_type="solid", fgColor="C6EFCE")   # 🟢 зелёный фон
RED_FILL      = PatternFill(fill_type="solid", fgColor="FFC7CE")   # 🔴 красный фон
GREEN_FONT    = Font(name="Arial", bold=True, size=10, color="006100")
RED_INDICATOR_FONT = Font(name="Arial", bold=True, size=10, color="9C0006")

THIN_SIDE   = Side(style="thin", color="BFBFBF")
THIN_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)

NUMBER_FORMAT = "#,##0.00"

# Столбцы листа "Договоры"
CONTRACT_COL_HEADERS = [
    "БИН",
    "Наименование компании",
    "Номер договора",
    "Краткое содержание",
    "Дата создания",
    "Сумма 1 (плановая/утвержденная, тг)",
    "Сумма 2 (фактически исполненная, тг)",
    "Общая итоговая сумма (тг)",
    "Максимальный доход (тг)",
    "Результат допустимого показателя загрузки",
    "Ссылка на договор",
]
CONTRACT_COL_WIDTHS = [16, 30, 28, 55, 22, 28, 28, 26, 26, 22, 50]

# Индексы числовых столбцов (1-based)
# F=6 (Сумма 1), G=7 (Сумма 2), H=8 (Итог), I=9 (Макс. доход), J=10 (Результат)
NUMERIC_COLS = {6, 7, 8, 9, 10}
LINK_COL = 11

# Столбцы листа "Сводка"
SUMMARY_COL_HEADERS = [
    "БИН компании",
    "Наименование компании",
    "Кол-во договоров",
    "Кол-во ошибок",
    "Общая итоговая сумма (тг)",
    "Максимальный доход (тг)",
    "Результат допустимого показателя загрузки",
    "Индикатор",
]
SUMMARY_COL_WIDTHS = [18, 30, 20, 16, 28, 26, 22, 14]

# Порог для индикатора
LOAD_THRESHOLD = 1.5


def _apply_header_style(cell, text: str) -> None:
    cell.value = text
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.border = THIN_BORDER
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _apply_cell_style(cell, value, *, is_number: bool = False, is_link: bool = False) -> None:
    cell.value = value
    cell.border = THIN_BORDER

    if is_link:
        cell.font = LINK_FONT
        cell.alignment = Alignment(vertical="top", wrap_text=False)
    elif is_number:
        cell.font = CELL_FONT
        cell.number_format = NUMBER_FORMAT
        cell.alignment = Alignment(horizontal="right", vertical="top")
    else:
        cell.font = CELL_FONT
        cell.alignment = Alignment(vertical="top", wrap_text=True)


# ── Лист «Договоры» ────────────────────────────────────────────────────────

def _write_contracts_sheet(wb: Workbook, results: list[ScrapeResult]) -> None:
    """Единый лист «Договоры» со всеми записями по всем БИН."""
    ws = wb.create_sheet(title="Договоры")

    # Заголовки
    ws.row_dimensions[1].height = 32
    for col_idx, (header, width) in enumerate(
        zip(CONTRACT_COL_HEADERS, CONTRACT_COL_WIDTHS), start=1
    ):
        cell = ws.cell(row=1, column=col_idx)
        _apply_header_style(cell, header)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    row_idx = 2
    for result in results:
        for record in result.records:
            ws.row_dimensions[row_idx].height = 15

            has_error = bool(record.error and record.error not in ("Сумма не найдена",))

            # A — БИН
            _apply_cell_style(ws.cell(row=row_idx, column=1), record.bin)

            # B — Наименование компании
            _apply_cell_style(ws.cell(row=row_idx, column=2),
                              record.supplier_name or "—")

            # C — Номер договора
            _apply_cell_style(ws.cell(row=row_idx, column=3),
                              record.contract_number or "—")

            # D — Краткое содержание
            desc = f"⚠ {record.error}" if has_error else record.description
            _apply_cell_style(ws.cell(row=row_idx, column=4), desc)

            # E — Дата создания
            _apply_cell_style(ws.cell(row=row_idx, column=5),
                              record.cr_datetime or "—")

            # F — Сумма 1 (плановая/утвержденная)
            cell_f = ws.cell(row=row_idx, column=6)
            if has_error:
                _apply_cell_style(cell_f, "—")
            else:
                _apply_cell_style(cell_f, record.amount_planned, is_number=True)

            # G — Сумма 2 (фактически исполненная)
            cell_g = ws.cell(row=row_idx, column=7)
            if has_error:
                _apply_cell_style(cell_g, "—")
            else:
                _apply_cell_style(cell_g, record.amount_actual, is_number=True)

            # H — Общая итоговая сумма (Сумма1 − Сумма2)
            cell_h = ws.cell(row=row_idx, column=8)
            if has_error:
                _apply_cell_style(cell_h, "—")
            else:
                _apply_cell_style(cell_h, record.amount_total, is_number=True)
                if record.amount_total < 0:
                    cell_h.font = RED_FONT

            # I — Максимальный доход
            cell_i = ws.cell(row=row_idx, column=9)
            if record.max_income > 0:
                _apply_cell_style(cell_i, record.max_income, is_number=True)
            else:
                _apply_cell_style(cell_i, "—")

            # J — Результат допустимого показателя загрузки
            cell_j = ws.cell(row=row_idx, column=10)
            if has_error or record.max_income <= 0:
                _apply_cell_style(cell_j, "—")
            else:
                load_ratio = record.amount_total / record.max_income
                _apply_cell_style(cell_j, load_ratio, is_number=True)

            # K — Ссылка на договор
            cell_k = ws.cell(row=row_idx, column=11)
            if record.url:
                cell_k.value = record.url
                cell_k.hyperlink = record.url
                cell_k.font = LINK_FONT
                cell_k.border = THIN_BORDER
                cell_k.alignment = Alignment(vertical="top", wrap_text=False)
            else:
                _apply_cell_style(cell_k, "—")

            row_idx += 1

    # Итоговая строка внизу таблицы
    ws.row_dimensions[row_idx].height = 20

    total_planned = sum(
        r.amount_planned for res in results for r in res.records
        if not r.error
    )
    total_actual = sum(
        r.amount_actual for res in results for r in res.records
        if not r.error
    )
    total_diff = total_planned - total_actual

    for col_idx in range(1, len(CONTRACT_COL_HEADERS) + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.fill = SUMMARY_FILL
        cell.border = THIN_BORDER
        cell.font = SUMMARY_FONT

    lbl = ws.cell(row=row_idx, column=4)
    lbl.value = "ИТОГО:"
    lbl.alignment = Alignment(horizontal="right", vertical="center")

    for col_idx, val in [(6, total_planned), (7, total_actual), (8, total_diff)]:
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.value = val
        cell.number_format = NUMBER_FORMAT
        cell.alignment = Alignment(horizontal="right", vertical="center")
        if col_idx == 8 and val < 0:
            cell.font = RED_BOLD_FONT

    ws.freeze_panes = "A2"


# ── Лист «Сводка» ──────────────────────────────────────────────────────────

def _write_summary_sheet(wb: Workbook, results: list[ScrapeResult]) -> None:
    """Лист «Сводка» — агрегированные итоги по каждому БИН + индикатор."""
    ws = wb.create_sheet(title="Сводка", index=0)

    # Заголовок
    ws.merge_cells("A1:H1")
    title_cell = ws.cell(row=1, column=1,
                         value="Сводный отчёт по государственным закупкам")
    title_cell.font = Font(name="Arial", bold=True, size=13)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # Дата формирования
    ws.merge_cells("A2:H2")
    date_cell = ws.cell(
        row=2, column=1,
        value=f"Сформировано: {datetime.now().strftime('%d.%m.%Y %H:%M:%S')}"
    )
    date_cell.font = Font(name="Arial", size=9, color="808080")
    date_cell.alignment = Alignment(horizontal="center")
    ws.row_dimensions[2].height = 18

    # Заголовки таблицы
    ws.row_dimensions[4].height = 28
    for col_idx, (header, width) in enumerate(
        zip(SUMMARY_COL_HEADERS, SUMMARY_COL_WIDTHS), start=1
    ):
        cell = ws.cell(row=4, column=col_idx)
        _apply_header_style(cell, header)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    row_idx = 5
    grand_total = 0.0
    grand_contracts = 0
    grand_errors = 0

    for result in results:
        valid_records = [r for r in result.records if not r.error]
        error_records = [r for r in result.records if r.error]
        bin_total = sum(r.amount_total for r in valid_records)
        contract_count = len(result.records)
        error_count = len(error_records)

        # Наименование компании — берём из первой записи, где оно есть
        supplier_name = ""
        for r in result.records:
            if r.supplier_name:
                supplier_name = r.supplier_name
                break

        grand_total += bin_total
        grand_contracts += contract_count
        grand_errors += error_count

        ws.row_dimensions[row_idx].height = 16

        def _sc(col, val, *, num=False, font=None, fill=None):
            c = ws.cell(row=row_idx, column=col, value=val)
            c.border = THIN_BORDER
            if font:
                c.font = font
            else:
                c.font = CELL_FONT
            if fill:
                c.fill = fill
            if num:
                c.number_format = NUMBER_FORMAT
                c.alignment = Alignment(horizontal="right", vertical="center")
            else:
                c.alignment = Alignment(horizontal="center", vertical="center")

        # A — БИН
        _sc(1, result.bin)

        # B — Наименование компании
        _sc(2, supplier_name or "—")

        # C — Кол-во договоров
        _sc(3, contract_count)

        # D — Кол-во ошибок
        _sc(4, error_count, font=RED_FONT if error_count > 0 else None)

        # E — Общая итоговая сумма
        _sc(5, bin_total, num=True)

        # F — Максимальный доход
        max_inc = result.max_income
        _sc(6, max_inc if max_inc > 0 else "—", num=max_inc > 0)

        # G — Результат допустимого показателя загрузки
        if max_inc > 0:
            load_ratio = bin_total / max_inc
            _sc(7, load_ratio, num=True)
        else:
            _sc(7, "—")

        # H — Индикатор
        indicator_cell = ws.cell(row=row_idx, column=8)
        indicator_cell.border = THIN_BORDER
        indicator_cell.alignment = Alignment(horizontal="center", vertical="center")
        if max_inc > 0:
            load_ratio = bin_total / max_inc
            if load_ratio > LOAD_THRESHOLD:
                # Красный кружок
                indicator_cell.value = "🔴"
                indicator_cell.fill = RED_FILL
                indicator_cell.font = RED_INDICATOR_FONT
            else:
                # Зелёный кружок
                indicator_cell.value = "🟢"
                indicator_cell.fill = GREEN_FILL
                indicator_cell.font = GREEN_FONT
        else:
            indicator_cell.value = "—"
            indicator_cell.font = CELL_FONT

        row_idx += 1

    # Итоговая строка
    total_row = row_idx
    ws.row_dimensions[total_row].height = 20

    for col_idx in range(1, len(SUMMARY_COL_HEADERS) + 1):
        c = ws.cell(row=total_row, column=col_idx)
        c.fill = TOTAL_FILL
        c.border = THIN_BORDER
        c.font = Font(name="Arial", bold=True, size=10)
        c.alignment = Alignment(horizontal="center", vertical="center")

    ws.cell(row=total_row, column=1).value = "ИТОГО"
    ws.cell(row=total_row, column=2).value = ""
    ws.cell(row=total_row, column=3).value = str(grand_contracts)
    ws.cell(row=total_row, column=4).value = str(grand_errors)

    cell_total = ws.cell(row=total_row, column=5)
    cell_total.value = grand_total
    cell_total.number_format = NUMBER_FORMAT
    cell_total.alignment = Alignment(horizontal="right", vertical="center")

    ws.cell(row=total_row, column=6).value = ""
    ws.cell(row=total_row, column=7).value = ""
    ws.cell(row=total_row, column=8).value = ""

    ws.freeze_panes = "A5"


# ── Публичный API ──────────────────────────────────────────────────────────

def build_excel_report(results: list[ScrapeResult]) -> bytes:
    """
    Формирует Excel-книгу по результатам парсинга.

    Структура:
      - Лист «Сводка» (первый) — агрегированные итоги по каждому БИН + индикатор
      - Лист «Договоры» — единая таблица всех договоров по всем БИН

    Возвращает байты .xlsx файла для st.download_button.
    """
    wb = Workbook()
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    _write_contracts_sheet(wb, results)
    _write_summary_sheet(wb, results)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def get_report_filename() -> str:
    return f"goszakup_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"