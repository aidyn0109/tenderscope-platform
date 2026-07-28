"""
excel_export_announcements.py — Формирование Excel-отчёта по результатам парсинга объявлений.

Новая структура (каждое объявление = одна строка):
  - БИН
  - Наименование компании
  - Номер объявления
  - Наименование объявления
  - Сумма 1 год (тг)
  - Ссылка на протокол
  - Ссылка на объявление
"""

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from scraper_announcements import ScrapeAnnouncementsResult, AnnouncementRecord

# ── Стили ──────────────────────────────────────────────────────────────────

HEADER_FONT  = Font(name="Arial", bold=True, size=10, color="000000")
CELL_FONT    = Font(name="Arial", size=10)
LINK_FONT    = Font(name="Arial", size=10, color="0563C1", underline="single")
SUMMARY_FONT = Font(name="Arial", bold=True, size=10)
HEADER_FILL  = PatternFill(fill_type="solid", fgColor="D9D9D9")
SUMMARY_FILL = PatternFill(fill_type="solid", fgColor="F2F2F2")

THIN_SIDE   = Side(style="thin", color="BFBFBF")
THIN_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)

NUMBER_FORMAT = "#,##0.00"

COL_HEADERS = [
    "БИН",
    "Наименование компании",
    "Номер объявления",
    "Наименование объявления",
    "Сумма 1 год (тг)",
    "Ссылка на протокол",
    "Ссылка на объявление",
]
COL_WIDTHS = [16, 35, 22, 55, 24, 50, 50]


def _apply_header_style(cell, text: str) -> None:
    cell.value = text
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.border = THIN_BORDER
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _apply_cell_style(cell, value, *, is_number: bool = False, is_link: bool = False) -> None:
    cell.value = value
    cell.border = THIN_BORDER
    if is_link:
        cell.font = LINK_FONT
        cell.alignment = Alignment(vertical="top", wrap_text=False)
    elif is_number:
        cell.font = CELL_FONT
        cell.number_format = NUMBER_FORMAT
        cell.alignment = Alignment(horizontal="right", vertical="top")
    else:
        cell.font = CELL_FONT
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def _write_announcements_sheet(wb: Workbook, result: ScrapeAnnouncementsResult) -> None:
    ws = wb.create_sheet(title="Объявления")

    # Заголовки
    ws.row_dimensions[1].height = 32
    for col_idx, (header, width) in enumerate(zip(COL_HEADERS, COL_WIDTHS), start=1):
        cell = ws.cell(row=1, column=col_idx)
        _apply_header_style(cell, header)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    row_idx = 2
    total_year1 = 0.0

    for record in result.results:
        ws.row_dimensions[row_idx].height = 15

        # A — БИН
        _apply_cell_style(ws.cell(row=row_idx, column=1), record.bin)

        # B — Наименование компании
        _apply_cell_style(ws.cell(row=row_idx, column=2), record.supplier_name or "—")

        # C — Номер объявления
        _apply_cell_style(ws.cell(row=row_idx, column=3), record.announcement_number or "—")

        # D — Наименование объявления
        _apply_cell_style(ws.cell(row=row_idx, column=4), record.announcement_name or "—")

        # E — Сумма 1 год
        cell_e = ws.cell(row=row_idx, column=5)
        if record.year1_sum > 0:
            _apply_cell_style(cell_e, record.year1_sum, is_number=True)
            total_year1 += record.year1_sum
        else:
            _apply_cell_style(cell_e, "—")

        # F — Ссылка на протокол
        cell_f = ws.cell(row=row_idx, column=6)
        if record.protocol_url:
            cell_f.value = record.protocol_url
            cell_f.hyperlink = record.protocol_url
            cell_f.font = LINK_FONT
            cell_f.border = THIN_BORDER
            cell_f.alignment = Alignment(vertical="top", wrap_text=False)
        else:
            _apply_cell_style(cell_f, "—")

        # G — Ссылка на объявление
        cell_g = ws.cell(row=row_idx, column=7)
        if record.announcement_url:
            cell_g.value = record.announcement_url
            cell_g.hyperlink = record.announcement_url
            cell_g.font = LINK_FONT
            cell_g.border = THIN_BORDER
            cell_g.alignment = Alignment(vertical="top", wrap_text=False)
        else:
            _apply_cell_style(cell_g, "—")

        row_idx += 1

    # Итоговая строка
    ws.row_dimensions[row_idx].height = 20
    for col_idx in range(1, len(COL_HEADERS) + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.fill = SUMMARY_FILL
        cell.border = THIN_BORDER
        cell.font = SUMMARY_FONT

    lbl = ws.cell(row=row_idx, column=4)
    lbl.value = "ИТОГО:"
    lbl.alignment = Alignment(horizontal="right", vertical="center")

    cell_total = ws.cell(row=row_idx, column=5)
    cell_total.value = total_year1
    cell_total.number_format = NUMBER_FORMAT
    cell_total.alignment = Alignment(horizontal="right", vertical="center")

    ws.freeze_panes = "A2"


def build_excel_announcements_report(result: ScrapeAnnouncementsResult) -> bytes:
    wb = Workbook()
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    _write_announcements_sheet(wb, result)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def get_announcements_report_filename() -> str:
    return f"goszakup_announcements_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"